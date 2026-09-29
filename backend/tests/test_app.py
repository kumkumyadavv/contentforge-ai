import json
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app.main import app
from app.services import brief as brief_module
from app.services.presentation import build_presentation
from app.services.validation import validate_output

client = TestClient(app)


@pytest.fixture(autouse=True)
def disable_external_model_calls(monkeypatch):
    monkeypatch.setattr('app.services.llm_client.settings', type('Settings', (), {'GEMINI_API_KEY': '', 'GEMINI_MODEL': 'test-model'})(), raising=False)


def test_health_endpoint():
    response = client.get('/health')
    assert response.status_code == 200
    assert response.json()['status'] == 'ok'


def test_source_validation_and_extraction():
    response = client.post('/api/source', data={'text': 'This is a test source with a date of 2026-01-15 and a key fact about revenue growth.'})
    assert response.status_code == 200
    payload = response.json()
    assert 'revenue growth' in payload['text']
    assert payload['metadata']['source_type'] == 'pasted_text'


def test_generate_returns_brief_and_outputs():
    response = client.post(
        '/api/generate',
        json={
            'source_text': 'Acme Ltd announced a 12% revenue increase in Q1 2026. The company plans to expand its product team. Risk includes shipping delays. The update is intended for executive leadership.',
            'config': {
                'audience': 'executive leaders',
                'tone': 'professional',
                'language': 'English',
                'detail_level': 'moderate',
                'communication_objective': 'inform and guide action',
                'output_types': ['executive_summary', 'advisory', 'linkedin'],
            },
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert set(payload) == {'brief', 'outputs', 'validation', 'source_metadata', 'generation_mode'}
    assert 'main_topic' in payload['brief']
    assert set(payload['outputs']) >= {'executive_summary', 'advisory', 'linkedin'}
    assert set(payload['validation']) >= {'executive_summary', 'advisory', 'linkedin'}
    assert payload['generation_mode'] in {'llm', 'fallback'}


@pytest.mark.parametrize('output_type', ['executive_summary', 'advisory', 'linkedin', 'x_post', 'infographic', 'presentation'])
def test_regenerate_uses_request_brief_and_validates_selected_output(output_type):
    app.state.current_brief = None
    brief = {
        'main_topic': 'Request supplied incident brief',
        'summary': 'The request brief is the source of truth.',
        'key_facts': ['The request brief is the source of truth.'],
        'dates': [],
        'entities': [],
        'impact': [],
        'risks': [],
        'recommended_actions': ['Review the request-supplied facts.'],
        'uncertainties': [],
        'evidence_references': ['Request brief reference.'],
    }

    response = client.post(
        '/api/output/regenerate',
        json={'output_type': output_type, 'config': {}, 'brief': brief},
    )

    assert response.status_code == 200
    payload = response.json()
    assert set(payload) == {'output_type', 'output', 'validation'}
    assert payload['output_type'] == output_type
    assert payload['validation']['status'] in {'passed', 'warning', 'failed'}
    if output_type in {'linkedin', 'x_post'}:
        assert 'Request supplied incident brief' in payload['output']['hook'] or 'request brief' in payload['output']['post'].lower()
    else:
        assert 'Request supplied incident brief' in payload['output']['title']
    assert app.state.current_brief is None


def test_regenerate_normalizes_output_type_before_validation():
    response = client.post(
        '/api/output/regenerate',
        json={
            'output_type': 'EXECUTIVE_SUMMARY',
            'config': {},
            'brief': {'main_topic': 'Uppercase output type'},
        },
    )

    assert response.status_code == 200
    assert response.json()['output_type'] == 'executive_summary'


def test_regenerate_rejects_unsupported_output_type():
    response = client.post(
        '/api/output/regenerate',
        json={
            'output_type': 'blog',
            'config': {},
            'brief': {'main_topic': 'Request brief'},
        },
    )

    assert response.status_code == 400
    assert response.json()['detail'] == 'Unsupported output type for regeneration.'


def test_new_outputs_include_brief_facts_dates_and_valid_x_character_count(monkeypatch):
    monkeypatch.setattr('app.services.llm_client.settings', type('Settings', (), {'GEMINI_API_KEY': '', 'GEMINI_MODEL': 'test-model'})(), raising=False)
    response = client.post(
        '/api/generate',
        json={
            'source_text': 'Acme Ltd reported an 18% increase on 2026-05-10. The team will review the impact and continue its investigation.',
            'config': {'output_types': ['x_post', 'infographic', 'presentation']},
        },
    )

    assert response.status_code == 200
    data = response.json()
    assert set(data['outputs']) == {'x_post', 'infographic', 'presentation'}
    x_post = data['outputs']['x_post']
    composed_post = '\n\n'.join([x_post['hook'], x_post['post'], x_post['call_to_action'], ' '.join(x_post['hashtags'])]).strip()
    assert x_post['character_count'] == len(composed_post)
    assert '18%' in json.dumps(data['outputs'])
    assert '2026-05-10' in json.dumps(data['outputs'])
    assert data['validation']['infographic']['status'] != 'failed'
    assert len(data['outputs']['presentation']['slides']) == 8


def test_uncertainty_validator_accepts_caveat_and_rejects_definitive_negative_claim():
    brief = {
        'summary': 'The investigation is ongoing.',
        'key_facts': [],
        'dates': [],
        'entities': [],
        'risks': ['No confirmed evidence that customer data was exfiltrated.'],
        'uncertainties': ['The investigation is ongoing.'],
    }
    valid = {
        'hook': 'Investigation update',
        'post': 'There was no confirmed evidence of customer data exfiltration, and the investigation remains ongoing.',
        'call_to_action': 'Follow for verified updates.',
        'hashtags': ['#Security'],
    }
    valid['character_count'] = len('\n\n'.join([valid['hook'], valid['post'], valid['call_to_action'], ' '.join(valid['hashtags'])]).strip())
    invalid = {**valid, 'post': 'Customer data was not exfiltrated.'}
    invalid['character_count'] = len('\n\n'.join([invalid['hook'], invalid['post'], invalid['call_to_action'], ' '.join(invalid['hashtags'])]).strip())

    assert validate_output('x_post', valid, brief).status != 'failed'
    failed = validate_output('x_post', invalid, brief)
    assert failed.status == 'failed'
    assert any('unconfirmed issue' in error.lower() for error in failed.errors)


def test_fallback_brief_preserves_no_confirmed_evidence_uncertainty(monkeypatch):
    monkeypatch.setattr('app.services.llm_client.settings', type('Settings', (), {'GEMINI_API_KEY': '', 'GEMINI_MODEL': 'test-model'})(), raising=False)
    source = 'There was no confirmed evidence that customer data was exfiltrated. The investigation is ongoing.'
    brief = brief_module.build_content_brief(source, {})
    outputs = brief_module.generate_output_variants(brief, {}, ['x_post'])

    assert any('no confirmed evidence' in item.lower() for item in brief['uncertainties'])
    assert 'ongoing' in outputs['x_post']['post'].lower()
    assert validate_output('x_post', outputs['x_post'], brief).status != 'failed'


def test_presentation_download_is_a_valid_pptx():
    output = {
        'title': 'Quarterly update',
        'subtitle': '18% increase in Q1 2026',
        'slides': [
            {'slide_number': index + 1, 'title': f'Slide {index + 1}', 'content': ['18% increase in Q1 2026'], 'speaker_notes': ''}
            for index in range(6)
        ],
    }
    content = build_presentation(output)
    presentation = Presentation(BytesIO(content))
    assert len(presentation.slides) == 6
    assert '18% increase in Q1 2026' in '\n'.join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)


def test_presentation_download_endpoint_streams_pptx():
    response = client.post(
        '/api/output/presentation/download',
        json={
            'output': {
                'title': 'Download test',
                'subtitle': 'Source grounded',
                'slides': [{'slide_number': index + 1, 'title': f'Slide {index + 1}', 'content': ['Verified source fact']} for index in range(6)],
            },
            'brief': {'summary': 'Verified source fact'},
        },
    )
    assert response.status_code == 200
    assert response.headers['content-type'].startswith('application/vnd.openxmlformats-officedocument.presentationml.presentation')
    assert len(Presentation(BytesIO(response.content)).slides) == 6


def test_brief_schema_has_expected_fields():
    response = client.post(
        '/api/generate',
        json={
            'source_text': 'The company is expanding into EU markets. It expects 30% faster onboarding in the next quarter.',
            'config': {'audience': 'team', 'tone': 'professional', 'language': 'English', 'detail_level': 'brief', 'communication_objective': 'inform', 'output_types': ['executive_summary']},
        },
    )
    assert response.status_code == 200
    brief = response.json()['brief']
    for field in ['main_topic', 'summary', 'key_facts', 'dates', 'entities', 'impact', 'risks', 'recommended_actions', 'uncertainties', 'evidence_references']:
        assert field in brief


def test_output_validation_checks_for_status_field():
    response = client.post(
        '/api/generate',
        json={
            'source_text': 'We are launching a training program in 2027 with a focus on retention and onboarding quality.',
            'config': {'audience': 'team', 'tone': 'professional', 'language': 'English', 'detail_level': 'brief', 'communication_objective': 'inform', 'output_types': ['linkedin']},
        },
    )
    data = response.json()['validation']['linkedin']
    assert 'status' in data
    assert data['status'] in {'passed', 'warning', 'failed'}


def test_deteministic_consistency_checks_are_returned():
    response = client.post(
        '/api/generate',
        json={
            'source_text': 'A 10% reduction in churn is expected by 2027 after the product launch. The team will continue to monitor customer feedback.',
            'config': {'audience': 'team', 'tone': 'professional', 'language': 'English', 'detail_level': 'brief', 'communication_objective': 'inform', 'output_types': ['executive_summary']},
        },
    )
    validation = response.json()['validation']['executive_summary']
    assert 'deterministic_checks' in validation
    assert len(validation['deterministic_checks']) > 0


def test_llm_generation_with_mocked_api_response(monkeypatch):
    class FakeModels:
        def generate_content(self, **kwargs):
            content = '{"main_topic": "Acme Ltd", "summary": "Acme Ltd reported 12% revenue growth in Q1 2026.", "key_facts": ["12% revenue growth in Q1 2026", "Expansion of product team"], "dates": ["Q1 2026"], "entities": ["Acme Ltd", "Q1 2026"], "impact": ["Potential operational scale-up"], "risks": ["Shipping delays may affect launch timing"], "recommended_actions": ["Prepare contingency for shipping"], "uncertainties": ["Shipping delays have not been confirmed."], "evidence_references": ["Acme source text"]}'
            return type('Response', (), {'text': content})()

    class FakeClient:
        def __init__(self, api_key):
            self.models = FakeModels()

    monkeypatch.setattr('app.services.llm_client.settings', type('Settings', (), {'GEMINI_API_KEY': 'test-key', 'GEMINI_MODEL': 'test-model'})(), raising=False)
    monkeypatch.setattr('app.services.llm_client.genai.Client', FakeClient)

    brief = brief_module.build_content_brief('Acme Ltd reported 12% revenue growth in Q1 2026. Shipping delays may affect launch timing.', {'audience': 'leadership'})
    assert brief['main_topic'] == 'Acme Ltd'
    assert '12% revenue growth' in brief['key_facts'][0]
    assert 'Shipping delays have not been confirmed.' in ' '.join(brief['uncertainties'])


def test_fallback_behavior_without_api_key(monkeypatch):
    monkeypatch.setattr('app.services.llm_client.settings', type('Settings', (), {'GEMINI_API_KEY': '', 'GEMINI_MODEL': 'test-model'})(), raising=False)
    brief = brief_module.build_content_brief('Operations were disrupted for 6 hours. Customer data exfiltration has not been confirmed.', {'audience': 'security team'})
    assert isinstance(brief, dict)
    assert 'Operations were disrupted for 6 hours.' in ' '.join(brief['key_facts']) or 'Operations were disrupted for 6 hours.' in brief['summary']


def test_uncertainty_and_fact_preservation_validation():
    brief = {
        'main_topic': 'Incident response',
        'summary': 'Operations were disrupted for 6 hours.',
        'key_facts': ['Operations were disrupted for 6 hours.', 'Customer data exfiltration has not been confirmed.'],
        'dates': ['2026-01-15'],
        'entities': ['Operations'],
        'impact': ['Operations were disrupted for 6 hours.'],
        'risks': ['Customer data exfiltration has not been confirmed.'],
        'recommended_actions': ['Validate the incident before making claims.'],
        'uncertainties': ['Customer data exfiltration has not been confirmed.'],
        'evidence_references': ['Source text'],
    }
    payload = {
        'title': 'Incident response',
        'one_line_summary': 'Operations were disrupted for 6 hours.',
        'situation': 'Operations were disrupted for 6 days.',
        'key_findings': ['Operations were disrupted for 6 days.'],
        'impact': ['Operations were disrupted for 6 days.'],
        'risks': ['Customer data was stolen.'],
        'recommended_actions': ['Notify stakeholders immediately.'],
        'uncertainties': ['Customer data exfiltration has not been confirmed.'],
        'source_references': ['Source text'],
    }
    validation = validate_output('executive_summary', payload, brief)
    assert validation.status == 'failed'
    assert any('unconfirmed' in item.lower() or 'mismatch' in item.lower() for item in validation.errors)


def test_ransomware_uncertainty_is_preserved_in_all_generated_outputs():
    brief = {
        'main_topic': 'Ransomware incident',
        'summary': 'Ransomware affected three servers. There is no confirmed evidence that customer payment information was exfiltrated; the investigation is ongoing.',
        'key_facts': ['Ransomware affected three servers.'],
        'dates': [],
        'entities': [],
        'impact': ['Three servers were affected.'],
        'risks': ['Customer payment information exfiltration is unconfirmed.'],
        'recommended_actions': ['Continue the investigation.'],
        'uncertainties': ['Customer payment information exfiltration is unconfirmed.'],
        'evidence_references': ['Ransomware incident source.'],
    }
    outputs = brief_module.generate_output_variants(brief, {})

    for output_type, payload in outputs.items():
        validation = validate_output(output_type, payload, brief)
        assert validation.status != 'failed', (output_type, validation.errors)

    confirmed_claim = {
        'title': 'Ransomware incident',
        'one_line_summary': 'Customer payment information was exfiltrated.',
        'situation': 'Customer payment information was exfiltrated.',
        'key_findings': ['Customer payment information was exfiltrated.'],
        'impact': ['Customer payment information was exfiltrated.'],
        'risks': ['Customer payment information was exfiltrated.'],
        'recommended_actions': ['Continue the investigation.'],
        'uncertainties': ['Customer payment information exfiltration is unconfirmed.'],
        'source_references': ['Ransomware incident source.'],
    }
    validation = validate_output('executive_summary', confirmed_claim, brief)
    assert validation.status == 'failed'
    assert any('unconfirmed issue' in item.lower() for item in validation.errors)
