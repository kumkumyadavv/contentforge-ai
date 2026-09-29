from __future__ import annotations

from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def build_presentation(output: dict[str, Any], brief: dict[str, Any] | None = None) -> bytes:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    for index, slide_data in enumerate(output.get("slides", [])):
        layout = presentation.slide_layouts[6]
        slide = presentation.slides.add_slide(layout)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(248, 249, 246)

        accent = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.16), Inches(7.5))
        accent.fill.solid()
        accent.fill.fore_color.rgb = RGBColor(25, 99, 88)
        accent.line.fill.background()

        if index == 0:
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.62), Inches(11.7), Inches(0.55))
            title_paragraph = title_box.text_frame.paragraphs[0]
            title_paragraph.text = str(output.get("title") or "Content Brief")
            title_paragraph.font.size = Pt(16)
            title_paragraph.font.bold = True
            title_paragraph.font.color.rgb = RGBColor(25, 99, 88)

            subtitle_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.7), Inches(1.2))
            subtitle_paragraph = subtitle_box.text_frame.paragraphs[0]
            subtitle_paragraph.text = str(output.get("subtitle") or "Evidence-grounded presentation")
            subtitle_paragraph.font.size = Pt(20)
            subtitle_paragraph.font.color.rgb = RGBColor(82, 94, 91)
            title_top = 2.6
        else:
            title_top = 0.8

        heading = slide.shapes.add_textbox(Inches(0.8), Inches(title_top), Inches(11.7), Inches(0.8))
        heading_paragraph = heading.text_frame.paragraphs[0]
        heading_paragraph.text = str(slide_data.get("title") or f"Slide {index + 1}")
        heading_paragraph.font.size = Pt(30)
        heading_paragraph.font.bold = True
        heading_paragraph.font.color.rgb = RGBColor(25, 43, 41)

        content = slide_data.get("content", [])
        body = slide.shapes.add_textbox(Inches(1.05), Inches(title_top + 1.05), Inches(11.2), Inches(4.5))
        frame = body.text_frame
        frame.word_wrap = True
        frame.margin_left = Inches(0.08)
        frame.margin_right = Inches(0.08)
        frame.clear()
        for bullet_index, item in enumerate(content[:5]):
            paragraph = frame.paragraphs[0] if bullet_index == 0 else frame.add_paragraph()
            paragraph.text = str(item)
            paragraph.level = 0
            paragraph.font.size = Pt(21)
            paragraph.font.color.rgb = RGBColor(45, 58, 55)
            paragraph.space_after = Pt(16)

        notes = slide.notes_slide.notes_text_frame
        note_text = str(slide_data.get("speaker_notes") or "")
        references = (brief or {}).get("evidence_references", [])
        if references:
            note_text = f"{note_text}\n\nSource references:\n" + "\n".join(str(item) for item in references)
        notes.text = note_text

        footer = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.7), Inches(0.25))
        footer_paragraph = footer.text_frame.paragraphs[0]
        footer_paragraph.text = f"CONTENTFORGE  /  {index + 1:02d}"
        footer_paragraph.font.size = Pt(9)
        footer_paragraph.font.color.rgb = RGBColor(108, 122, 117)
        footer_paragraph.alignment = PP_ALIGN.RIGHT

    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()